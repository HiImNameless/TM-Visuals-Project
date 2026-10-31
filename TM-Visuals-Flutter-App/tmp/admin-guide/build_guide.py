from pathlib import Path
from datetime import date
import json, re, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/docx/TM_Visuals_Admin_Features_Code_Guide.docx'
SOURCES={p.relative_to(ROOT/'lib').as_posix():p.read_text(encoding='utf-8-sig') for p in (ROOT/'lib').rglob('*.dart')}
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    if st.element.rPr is not None:
        for c in list(st.element.rPr):
            if c.tag==qn('w:color'):
                for attr in [qn('w:themeColor'),qn('w:themeTint'),qn('w:themeShade')]:c.attrib.pop(attr,None)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(27)
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next=True
head=sec.header.paragraphs[0];head.text='TM VISUALS   |   ADMIN FEATURE CODE GUIDE'
for r in head.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor(0,0,0)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('31 October 2026   |   Page ').font.size=Pt(8.5)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
pages=[];entries=[]

def p(text,boldlead=None):
    par=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):par.add_run(boldlead).bold=True;par.add_run(text[len(boldlead):])
    else:par.add_run(text)
    return par
def code(text):
    par=doc.add_paragraph();par.paragraph_format.space_after=Pt(5)
    run=par.add_run(text);run.font.name='Consolas';run.font.size=Pt(9)
    return par
def page(title,intro=None):
    if pages:doc.add_page_break()
    pages.append(title);doc.add_heading(f'{len(pages)}  {title}',level=1)
    if intro:p(intro)
def h(title,text=None):
    doc.add_heading(title,level=2)
    if text:p(text)
def say(text):
    par=p('How to explain it: '+text)
    par.runs[0].italic=True
def source(file,needle,after=None):
    text=SOURCES[file];start=text.find(after) if after else 0
    assert start>=0,(file,after)
    ix=text.find(needle,start);assert ix>=0,(file,needle)
    return text.count('\n',0,ix)+1
def fn(title,sig,body,file,needle,after=None):
    h(title);line=source(file,needle,after)
    code(f'{sig}\nlib/{file}  |  line {line}')
    p(body);entries.append({'file':file,'symbol':sig,'line':line,'chapter':len(pages)})
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):cells[i].text=str(val)
    for rid,row in enumerate(t.rows):
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                n=OxmlElement(f'w:{side}');n.set(qn('w:w'),'85');n.set(qn('w:type'),'dxa');margins.append(n)
            pr.append(margins);b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                n=OxmlElement(f'w:{side}');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'D9D9D9');b.append(n)
            pr.append(b)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243342' if rid==0 else ('F2F5F7' if rid%2==0 else 'FFFFFF'));pr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(2)
                for run in par.runs:
                    run.font.size=Pt(10)
                    if rid==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

page('Purpose and presentation overview')
doc.add_paragraph('TM Visuals admin features',style='Title')
p('A study and presentation guide for Brandon\'s admin calendar, whole-day availability manager and existing invoice PDF upload. The guide explains what each relevant file and function does, how data moves through the application, and what to demonstrate.')
p('The central design is a separation of responsibilities: views display the interface, viewmodels manage screen state, models describe data, and services communicate with Supabase. main.dart is only the startup entry point.')
h('A short opening you can use')
say('My assigned features are the admin calendar, date availability management and uploading existing invoice PDFs. I used the team\'s MVVM structure with Provider, so the screens are separate from state management and database operations. An administrator can inspect bookings, close or reopen an entire date without deleting bookings, and attach a PDF invoice to a booking.')
h('What this guide covers')
table(['Feature','What to explain'],[
('Calendar','Month navigation, booking previews, date colours and loading data.'),
('Date manager','One availability setting for the whole date and preservation of bookings.'),
('Invoice upload','Validation, private file storage, invoice metadata and safe retries.'),
('Supporting code','Admin sign-in, protected routes, shared widgets and bottom navigation.')],[1.35,5.65])
h('How to study')
p('Read chapters 2 and 3 first, then follow the calendar, availability and invoice chapters in order. Use chapters 22 and 23 to rehearse the demo and questions. In Word, search for a function name with Ctrl+F or open the Navigation Pane to jump between headings.')
p('All source paths are relative to the Flutter app folder. Line references were checked against the working tree on 31 October 2026, based on commit 396959b. They may move after later edits. This guide describes current source code; it is not a fresh certification of the hosted database or a live upload test.')

page('Architecture and the flow of data','MVVM means Model View ViewModel. The service layer is an additional boundary for network and Storage work.')
table(['Layer','Role','Example'],[
('View','Builds widgets and forwards user actions.','views/calendar.dart'),
('ViewModel','Stores state and coordinates operations.','viewmodels/calendar_viewmodel.dart'),
('Model','Represents typed application data.','models/calendar_model.dart'),
('Service','Reads and writes external data.','services/calendar_service.dart'),
('Supabase','Provides Auth, Postgres and Storage.','time_slot, date_availability, invoice')],[1,3.2,2.8])
h('A calendar request from start to finish')
p('CalendarScreen creates CalendarViewModel with a CalendarDataSource. The viewmodel calls loadMonth and loadAvailability. CalendarService queries Supabase, then CalendarSlot.fromMap converts database maps into objects. The viewmodel saves those results and notifies its listeners. The view rebuilds and displays the returned state.')
code('User action -> View -> ViewModel -> Service -> Supabase\nSupabase result -> Service -> Model objects -> ViewModel -> View')
h('How Provider updates the screen')
p('ChangeNotifier is the base class that lets a viewmodel announce state changes. notifyListeners tells listening widgets that something changed. context.watch<T>() subscribes a widget to that state. context.read<T>() obtains an object without subscribing the caller. Consumer<T> rebuilds a selected part of the UI.')
h('Why services are injectable')
p('CalendarDataSource and InvoiceDataSource define contracts. A viewmodel can receive a real Supabase implementation or a fake implementation for a test. This makes it possible to test state changes and error handling without connecting to the live project.')
h('Why the files are separate')
p('A colour or layout change belongs in a view or shared widget. A database query change belongs in a service. A rule about when saving is allowed belongs in a viewmodel. A field or database mapping change belongs in a model. This separation makes code easier to explain, review and test.')
say('Provider connects the screen to its viewmodel. The service retrieves data, the model gives it a clear structure, and the viewmodel tells the screen when to rebuild.')

page('Startup configuration and application setup')
fn('Start the Flutter application','main() -> Future<void>', 'Initializes the Flutter binding, reads SUPABASE_URL and SUPABASE_PUBLISHABLE_KEY from Dart compile-time defines, and creates a Supabase client if usable values exist. It rejects a key beginning sb_secret_. Initialization failures are captured in error. runApp receives MyApp with either a client or a startup error. No calendar, date-editor or invoice form layout is built here.','main.dart','Future<void> main()')
fn('Provide shared application state','MyApp constructor and build(context) -> Widget', 'Accepts an optional client, startupError and injectable authService. MultiProvider registers the team\'s shared viewmodels. MaterialApp sets the theme and public root route. With no client it shows SupabaseSetupScreen; with a client it shows AdminEntryScreen. Current team integration registers BookingsViewModel using BookingListViewModel. Those booking classes are integration dependencies, not Brandon\'s calendar implementation.','app.dart','Widget build(')
fn('Explain missing connection settings','SupabaseSetupScreen.build(context) -> Widget', 'Displays setup instructions when settings are missing or initialization failed. The message is selected using the optional error value. This screen does not connect to Supabase or write configuration. The example file mentioned in its text is currently deleted in the working tree, so check the existing local configuration rather than assuming that example is present.','views/supabase_setup.dart','Widget build(')
h('How the local configuration is used')
code('flutter run --dart-define-from-file=config/supabase.local.json')
p('The build command supplies values from JSON to String.fromEnvironment. The app does not read that JSON dynamically while running. Restart after changing it. Never display the key or the local file contents during a presentation. A public publishable key is intended for a client application, but database access must still be restricted by server-side policies.')
say('main.dart prepares the connection and starts the app. MyApp configures providers, styling and the first screen. The actual feature screens live under views.')

page('Authentication service and sign in screen')
fn('Create the authentication service','AdminAuthService(client)', 'Stores the SupabaseClient that all authentication operations use. Supplying the client allows tests or different app hosts to provide the connection rather than creating it inside every function.','services/admin_auth_service.dart','AdminAuthService(this.client)')
fn('Observe the current user and session changes','changes and userId getters', 'changes exposes client.auth.onAuthStateChange. userId returns the current Auth user UUID or null when signed out. The UUID is used to look up administrator membership. An Auth identity and an application administrator record are two separate things.','services/admin_auth_service.dart','Stream<AuthState> get changes')
fn('Look up administrator membership','currentAdminId() -> Future<String?>', 'Returns null if no user is signed in. Otherwise queries the current source\'s admins table for admin_id where user_id equals the Auth UUID. maybeSingle allows either one row or no row. A matching row produces an admin ID; no row means the account is not recognized as an administrator. Confirm the live table name before your demo.','services/admin_auth_service.dart','Future<String?> currentAdminId()')
fn('Sign in and sign out','signIn(email, password) and signOut()', 'signIn trims the email and calls signInWithPassword; it leaves the password unchanged. signOut delegates to Supabase Auth. The service does not render errors or navigate. The viewmodel responds to the resulting session events.','services/admin_auth_service.dart','Future<void> signIn(')
fn('Collect and validate sign in input','AdminSignInScreen.createState and state build(context)', 'createState creates the screen\'s state object. A GlobalKey<FormState> gives the local submit callback access to validators. Two TextEditingControllers hold the input. The email validator checks for @; the password validator checks that it is not empty. submit calls vm.signIn only if the form validates. Password input is obscured, and vm.busy disables the form and button during a request. vm.error is displayed below the inputs.','views/admin_sign_in.dart','Widget build(',after='class _AdminSignInScreenState')
fn('Release input controllers','_AdminSignInScreenState.dispose()', 'Disposes the email and password controllers and then calls super.dispose. This releases resources when the sign-in page is removed. It does not sign the user out.','views/admin_sign_in.dart','void dispose()')

page('Authentication viewmodel and session state')
p('AdminAuthViewModel owns userId, adminId, error, checking and busy. authorized is true only when both IDs are non-null. _version identifies the newest membership check, and _disposed prevents notifications after the viewmodel is removed.')
fn('Start listening for authentication events','AdminAuthViewModel(service)', 'Subscribes to service.changes and calls check immediately. Each normal event starts another membership check. A stream error invalidates older checks, clears adminId and exposes the error instead of leaving the admin workspace authorized.','viewmodels/admin_auth_viewmodel.dart','AdminAuthViewModel(this.service)')
fn('Determine whether access is authorized','authorized getter', 'Returns userId != null && adminId != null. This is the UI gate condition. It is not a replacement for Supabase Row Level Security or Storage policies.','viewmodels/admin_auth_viewmodel.dart','bool get authorized')
fn('Recheck identity and administrator access','check() -> Future<void>', 'Increments _version, gets the current user, and decides whether an already authorized same-user navigator can remain visible while the check runs. A signed-out user needs no membership query. A signed-in user triggers currentAdminId. Only the newest result may update state. Missing membership or a query error removes authorization; finally clears checking and updates listeners.','viewmodels/admin_auth_viewmodel.dart','Future<void> check()')
fn('Run sign in and sign out operations','signIn(email, password) and signOut()', 'The busy guard blocks duplicate actions. Each method marks itself busy, calls the authentication service, captures an error for the UI and clears busy in finally. Authentication events, rather than the button itself, drive the subsequent membership check and screen selection.','viewmodels/admin_auth_viewmodel.dart','Future<void> signIn(')
fn('Notify safely and release the subscription','_changed() and dispose()', '_changed calls notifyListeners only while the object is alive. dispose marks it disposed, increments _version so pending checks become stale, cancels the subscription and calls the superclass cleanup.','viewmodels/admin_auth_viewmodel.dart','void _changed()')
say('Logging in proves who the user is. The membership check then decides whether that user may enter the administrator area. The version counter stops an older request from restoring access after a newer session change.')

page('The admin gate protected routes and draft lifetime')
fn('Create authentication state for the entry screen','AdminEntryScreen.build(context)', 'Creates AdminAuthViewModel inside ChangeNotifierProvider, using the injected auth service or the real AdminAuthService. Its child is _AdminGate, which listens to that viewmodel.','views/admin_entry.dart','Widget build(',after='class AdminEntryScreen')
fn('Choose the correct screen and protect the route stack','_AdminGate.build(context)', 'Shows sign-in when userId is null, a spinner during the initial membership check, or an access/error page when authorization is missing. For an authorized user, creates an AddInvoiceViewModel provider and a nested Navigator. The navigator starts at /, which maps to Calendar. Removing this authorized subtree on sign-out removes open protected routes and the in-memory invoice draft.','views/admin_entry.dart','Widget build(',after='class _AdminGate')
fn('Build a named route','RouteManager.getRoute(settings, client) -> Route', 'Selects a WidgetBuilder by settings.name and returns MaterialPageRoute with the original settings. A DateTime argument selects the date editor. A String argument preselects a booking for invoice upload. Add invoice reads the shared draft viewmodel from Provider. Unknown names produce a page-not-found screen. The router belongs to the authenticated nested Navigator.','routes/route_manager.dart','static Route<dynamic> getRoute(')
table(['Route','Destination'],[('/ and /calendar','CalendarScreen'),('/date-availability','DateManagerScreen with a DateTime argument'),('/add-invoice','AddInvoiceScreen with an optional booking ID'),('/bookings and /admin-profile','Existing team screens'),('/more','AdminMoreScreen'),('/invoices /services /settings','Existing team screens'),('/portfolio-content /portfolio-gallery /clients','Current team integrations')],[2.65,4.35])
h('Why the draft provider is above the navigator')
p('The upload form can close while its viewmodel remains alive in the authorized subtree. Reopening the form can therefore resume an uncertain upload with the same ID and bytes. The provider key includes the user ID. This is temporary in-memory recovery, not a draft saved to a database or disk.')
